#!/usr/bin/env python3
import os
import sys
import glob
import re
import argparse
import subprocess
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def print_help():
    help_text = """
Usage: python nmapp.py [OPTIONS] [NMAP_FILES...]

Options:
  -h, --help            Show this help message and syntax
  -iL <target_file>     Scan a list of targets from a file using the automated nmapp workflow
  -o, --output <file>   Output Excel filename (default: nmap_port_summary.xlsx)

Examples:
  python nmapp.py open-services-*.nmap
  python nmapp.py -iL targets.txt
  python nmapp.py -iL targets.txt -o custom_report.xlsx
"""
    print(help_text)

def run_nmap_workflow(target):
    print(f"\033[34m[*] Running Fast TCP Nmap Scan on {target}\033[0m")
    fast_out = f"fast-tcp-{target}.nmap"
    subprocess.run(f"nmap -F -T4 -Pn {target} -oN {fast_out}", shell=True, capture_output=True, text=True)
    
    open_fast_ports = []
    if os.path.exists(fast_out):
        with open(fast_out, 'r') as f:
            for line in f:
                if 'open' in line and '/tcp' in line:
                    p = line.split('/')[0].strip()
                    if p.isdigit():
                        open_fast_ports.append(p)
    fnmap = ",".join(open_fast_ports)
    print(f"\033[36m[+] Opened Ports for {target} - {fnmap}\033[0m")
    
    print(f"\033[34m[*] Running Full TCP Nmap Scan on {target}\033[0m")
    all_out = f"all-tcp-{target}.nmap"
    subprocess.run(f"nmap -p- -T4 -Pn {target} -oN {all_out}", shell=True, capture_output=True, text=True)
    
    open_all_ports = []
    if os.path.exists(all_out):
        with open(all_out, 'r') as f:
            for line in f:
                if 'open' in line and '/tcp' in line:
                    p = line.split('/')[0].strip()
                    if p.isdigit():
                        open_all_ports.append(p)
    qnmap = ",".join(open_all_ports)
    print(f"\033[36m[+] Opened Ports for {target} - {qnmap}\033[0m")
    
    print(f"\033[34m[*] Running Open Service Nmap Scan on {target}\033[0m")
    services_out = f"open-services-{target}.nmap"
    if qnmap:
        subprocess.run(f"nmap -p {qnmap} -sCV {target} -oN {services_out}", shell=True)
    else:
        subprocess.run(f"nmap -sCV {target} -oN {services_out}", shell=True)
    print(f"\033[36m[+] Finished Scanning on {target}\033[0m")
    return services_out

def parse_nmap_on_file(file_path):
    hosts = []
    if not os.path.exists(file_path):
        print(f"Warning: File {file_path} not found.")
        return hosts
        
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        content = f.read()
        
    current_ip = None
    current_hostname = ""
    open_ports_map = {}
    
    lines = content.splitlines()
    for line in lines:
        if line.startswith("Nmap scan report for"):
            if current_ip:
                hosts.append({'ip': current_ip, 'hostname': current_hostname, 'open_ports_map': open_ports_map})
                open_ports_map = {}
                current_hostname = ""
                current_ip = None
                
            parts = line.replace("Nmap scan report for", "").strip()
            if "(" in parts and ")" in parts:
                match = re.search(r'^(.*?)\s+\(([0-9]+\.[0-9]+\.[0-9]+\.[0-9]+)\)$', parts)
                if match:
                    current_hostname = match.group(1).strip()
                    current_ip = match.group(2).strip()
            else:
                current_ip = parts.strip()
                current_hostname = ""
                
        port_match = re.match(r'^([0-9]+)/(tcp|udp)\s+open\s+([\w\-\?]+)', line)
        if port_match and current_ip:
            try:
                port_num = int(port_match.group(1))
                service_name = port_match.group(3).strip()
                open_ports_map[port_num] = service_name
            except ValueError:
                pass
                
    if current_ip:
        hosts.append({'ip': current_ip, 'hostname': current_hostname, 'open_ports_map': open_ports_map})
        
    return hosts

def generate_nmap_excel(file_patterns, output_filename="nmap_port_summary.xlsx"):
    fixed_ports = [21, 22, 23, 25, 53, 80, 88, 135, 139, 443, 445, 1433, 3306, 3389, 5985, 8000, 8080, 8443]
    default_fixed_services = {
        21: "ftp", 22: "ssh", 23: "telnet", 25: "smtp", 53: "domain", 80: "http", 88: "kerberos-sec",
        135: "msrpc", 139: "netbios-ssn", 443: "https", 445: "microsoft-ds", 1433: "ms-sql-s",
        3306: "mysql", 3389: "ms-wbt-server", 5985: "http", 8000: "http-alt", 8080: "http-proxy", 8443: "https-alt"
    }
    
    port_to_service = default_fixed_services.copy()
    all_seen_ports = set(fixed_ports)
    all_parsed_hosts = []
    
    file_paths = []
    for pat in file_patterns:
        matched = glob.glob(pat)
        if matched:
            file_paths.extend(matched)
        elif os.path.exists(pat):
            file_paths.append(pat)
            
    if not file_paths:
        print("Error: No valid Nmap files found to process.", file=sys.stderr)
        print_help()
        sys.exit(1)
        
    for fp in file_paths:
        hosts = parse_nmap_on_file(fp)
        for h in hosts:
            for p, svc in h['open_ports_map'].items():
                all_seen_ports.add(p)
                if p not in port_to_service or port_to_service[p] == "unknown":
                    port_to_service[p] = svc
            all_parsed_hosts.append(h)
            
    other_ports = sorted([p for p in all_seen_ports if p not in fixed_ports])
    all_columns_ports = fixed_ports + other_ports
    
    port_headers = {}
    for p in all_columns_ports:
        svc = port_to_service.get(p, "unknown").strip()
        if not svc:
            svc = "unknown"
        port_headers[str(p)] = f"{p}\n({svc})"
        
    rows = []
    for h in all_parsed_hosts:
        row = {
            'IP': h['ip'],
            'DNS/Hostname': h['hostname']
        }
        for p in all_columns_ports:
            p_str = str(p)
            row[p_str] = "✓" if p in h['open_ports_map'] else ""
        rows.append(row)
        
    df = pd.DataFrame(rows)
    rename_map = {str(p): port_headers[str(p)] for p in all_columns_ports}
    df = df.rename(columns=rename_map)
    
    cols_order = ['IP', 'DNS/Hostname'] + [port_headers[str(p)] for p in all_columns_ports]
    
    if not df.empty:
        for c in cols_order:
            if c not in df.columns:
                df[c] = ""
        df = df[cols_order]
        df = df.drop_duplicates(subset=['IP'], keep='last')
    else:
        df = pd.DataFrame(columns=cols_order)
        
    wb = Workbook()
    ws = wb.active
    ws.title = "Port Summary"
    
    headers = list(df.columns)
    ws.append(headers)
    
    for _, row in df.iterrows():
        ws.append(list(row))
        
    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    ticked_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    ticked_font = Font(name="Arial", size=10, bold=True, color="006100")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    cell_font = Font(name="Arial", size=10)
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border
        
    for row_num in range(2, ws.max_row + 1):
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.border = thin_border
            
            if cell.value == "✓":
                cell.fill = ticked_fill
                cell.font = ticked_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.font = cell_font
                if col_num == 1 or col_num > 2:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                val_str = str(cell.value)
                lines = val_str.split('\n')
                max_len = max(max_len, max(len(l) for l in lines))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
        
    ws.row_dimensions[1].height = 35
    
    wb.save(output_filename)
    print(f"\033[32m[+] Excel report successfully generated and saved as {output_filename}\033[0m")

def main():
    if len(sys.argv) < 2 or sys.argv[1] in ['-h', '--help']:
        print_help()
        sys.exit(0 if len(sys.argv) > 1 else 1)
        
    output_file = "nmap_port_summary.xlsx"
    targets_to_scan = []
    file_args = []
    
    i = 1
    while i < len(sys.argv):
        arg = sys.argv[i]
        if arg in ['-h', '--help']:
            print_help()
            sys.exit(0)
        elif arg == '-iL':
            if i + 1 < len(sys.argv):
                target_list_file = sys.argv[i+1]
                if os.path.exists(target_list_file):
                    with open(target_list_file, 'r') as tf:
                        for line in tf:
                            t = line.strip()
                            if t and not t.startswith('#'):
                                targets_to_scan.append(t)
                else:
                    print(f"Error: Target list file '{target_list_file}' not found.", file=sys.stderr)
                    sys.exit(1)
                i += 2
            else:
                print("Error: -iL requires a target file argument.", file=sys.stderr)
                sys.exit(1)
        elif arg in ['-o', '--output']:
            if i + 1 < len(sys.argv):
                output_file = sys.argv[i+1]
                i += 2
            else:
                print("Error: Output flag requires a filename.", file=sys.stderr)
                sys.exit(1)
        else:
            file_args.append(arg)
            i += 1
            
    generated_files = []
    if targets_to_scan:
        total_targets = len(targets_to_scan)
        for idx, target in enumerate(targets_to_scan, 1):
            print(f"\n\033[1m[Processing Target {idx}/{total_targets}]\033[0m")
            svc_file = run_nmap_workflow(target)
            if os.path.exists(svc_file):
                generated_files.append(svc_file)
    else:
        generated_files = file_args
        
    if not generated_files:
        print("Error: No Nmap files or targets specified.", file=sys.stderr)
        print_help()
        sys.exit(1)
        
    generate_nmap_excel(generated_files, output_file)

if __name__ == '__main__':
    main()
